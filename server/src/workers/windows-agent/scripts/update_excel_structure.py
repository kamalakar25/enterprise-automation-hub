import argparse
import csv
import datetime as dt
import os
import re
import sys
import shutil
from collections import Counter
from pathlib import Path

import openpyxl

BASE_DIR = r"E:\Kamalakar\Projects\SAP daily updated"
DEFAULT_XLSX = os.path.join(BASE_DIR, "SAP_Daily_Input.xlsx")
SHEET_NAME = "DAILY_INPUT"
HEADER_ROW = 2
DATA_START_ROW = HEADER_ROW + 1

VALID_CUMMODES = {"EMAIL", "PHONE", "VISIT", "SANYOG"}
DONE_TEXT = "Updated Successfully"

# Excel automation rules
SUPPLIER_LEAD_TIME_HEADER = "Supplier Lead Time"
COMMITTED_DATE_HEADER = "Supplier Committed Date"
AUTOMATION_MODE_HEADER = "Automation Mode"
AUTO_REMARK_TEMPLATE = "Material is expected to dispatch by {date}"
PULL_IN_REMARK_TEMPLATE = "Request supplier to pull in the committed date of {date}."
# Rows with these DFT/status meanings are intentionally left for manual handling.
MANUAL_DFT_CODES = {"10", "11", "27"}  # Ready for Dispatch, In Transit, Partial Qty In Transit
MANUAL_REMARK_KEYWORDS = (
    "partial dispatch",
    "partial dispatched",
    "partial qty",
    "in transit",
    "ready for dispatch",
    "no committed date",
    "no commitment",
)


# Expected physical layout: A:S / 19 columns.
COLS = {
    "system_status": 1,
    "vendor_code": 2,
    "vendor_name": 3,
    "po": 4,
    "po_date": 5,
    "line": 6,
    "material_code": 7,
    "sched_date": 8,
    "qty": 9,
    "material_desc": 10,
    "exp_date": 11,
    "remarks": 12,
    "dft": 13,
    "status_in": 14,
    "pernr": 15,
    "cummode": 16,
    "status_out": 17,
    "err_msg": 18,
    "upd_time": 19,
}

HEADER_CHECKS = [
    (4, "PO"),
    (6, "Line"),
    (7, "Material"),
    (8, "Scheduled"),
    (9, "Pending"),
    (10, "Material"),
    (11, "Expected"),
    (12, "Remarks"),
    (13, "DFT"),
    (14, "Status"),
    (15, "Emp"),
    (16, "Communication"),
    (17, "Update"),
    (18, "Error"),
    (19, "Last"),
]

WORKLIST_FIELDS = [
    "RunID",
    "ExcelRow",
    "VendorName",
    "PO",
    "Line",
    "MaterialCode",
    "SchedDate",
    "Qty",
    "MatchKey",
    "KeyOccurrence",
    "KeyTotal",
    "SAPRow",
    "ExpDate",
    "Remarks",
    "DFT",
    "PERNR",
    "CUMMODE",
    "StatusIn",
    "StatusOut",
]

STATUS_FIELDS = [
    "RunID",
    "ExcelRow",
    "VendorName",
    "PO",
    "Line",
    "MaterialCode",
    "SchedDate",
    "Qty",
    "PERNR",
    "CUMMODE",
    "Status",
    "ErrorRemarks",
    "Timestamp",
]


# Pre-compiled regular expressions for high-performance parsing
RE_NORM_SPACE = re.compile(r"\s+")
RE_DATE_MATCH = re.compile(r"(\d{1,2}[./]\d{1,2}[./]\d{2,4}|\d{4}[-/]\d{1,2}[-/]\d{1,2})")
RE_PURE_PATTERNS = (
    re.compile(r"^\d{1,2}\.\d{1,2}\.(\d{2}|\d{4})$"),
    re.compile(r"^\d{1,2}/\d{1,2}/(\d{2}|\d{4})$"),
    re.compile(r"^\d{4}-\d{1,2}-\d{1,2}$"),
    re.compile(r"^\d{4}/\d{1,2}/\d{1,2}$"),
)
RE_DATE_PATTERNS = (
    (re.compile(r"^(\d{1,2})\.(\d{1,2})\.(\d{2}|\d{4})$"), "dmy"),
    (re.compile(r"^(\d{1,2})/(\d{1,2})/(\d{2}|\d{4})$"), "dmy"),
    (re.compile(r"^(\d{4})-(\d{1,2})-(\d{1,2})$"), "ymd"),
    (re.compile(r"^(\d{4})/\d{1,2}/\d{1,2}$"), "ymd"),
)
RE_DATE_FMT_DMY_DOT = re.compile(r"^(\d{1,2})\.(\d{1,2})\.(\d{4})$")
RE_DATE_FMT_YMD = re.compile(r"^(\d{4})[-/](\d{1,2})[-/](\d{1,2})$")
RE_DATE_FMT_DMY_SLASH = re.compile(r"^(\d{1,2})/(\d{1,2})/(\d{4})$")

RE_NON_DIGITS = re.compile(r"\D")
RE_QTY_CLEAN = re.compile(r"[^0-9,\.\-]")
RE_DFT_CODE = re.compile(r"^(\d+)")


def sanitize(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).replace("\t", " ").replace("\r", " ").strip()


def norm_header(value) -> str:
    return RE_NORM_SPACE.sub(" ", sanitize(value)).lower()


def header_contains(value, *needles) -> bool:
    h = norm_header(value)
    return all(n.lower() in h for n in needles)


def find_header_col(ws, *needles):
    for col in range(1, ws.max_column + 1):
        if header_contains(ws.cell(HEADER_ROW, col).value, *needles):
            return col
    return None


def ensure_header_col(ws, header_name):
    existing = find_header_col(ws, *header_name.split())
    if existing:
        return existing, False
    col = ws.max_column + 1
    ws.cell(HEADER_ROW, col).value = header_name
    return col, True


def parse_date_only(value):
    """Parse only when the entire cell is just a date, with no extra remark text."""
    if value is None or value == "":
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    text = sanitize(value).strip().rstrip(".")
    if not text:
        return None
    if not any(p.fullmatch(text) for p in RE_PURE_PATTERNS):
        return None
    return parse_date(text)


def is_auto_generated_remark(text):
    t = sanitize(text).lower()
    return t.startswith("material is expected to dispatch by") or t.startswith("request supplier to pull in")


def parse_date(value):
    """Parse common Calc/Excel/SAP date values to datetime.date; return None if not parseable."""
    if value is None or value == "":
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value

    text = sanitize(value)
    if not text:
        return None

    # Extract date even from generated remarks such as:
    # "Material is expected to dispatch by 15.06.2026" or
    # "Request supplier to pull in the committed date of 15/06/2026."
    date_match = RE_DATE_MATCH.search(text)
    if date_match:
        text = date_match.group(1)
    else:
        text = text.split()[0]
    text = text.rstrip(".")

    for pattern, order in RE_DATE_PATTERNS:
        m = pattern.fullmatch(text)
        if not m:
            continue
        a, b, c = m.groups()
        try:
            if order == "dmy":
                day, month, year = int(a), int(b), int(c)
            else:
                year, month, day = int(a), int(b), int(c)
            if year < 100:
                year += 2000
            return dt.date(year, month, day)
        except ValueError:
            return None
    return None


def format_date(value) -> str:
    """Return dd.mm.yyyy for SAP/VBS, safely handling Excel dates and strings."""
    if value is None or value == "":
        return ""
    if isinstance(value, dt.datetime):
        value = value.date()
    if isinstance(value, dt.date):
        return value.strftime("%d.%m.%Y")

    parsed = parse_date(value)
    if parsed:
        return parsed.strftime("%d.%m.%Y")

    text = sanitize(value)
    if not text:
        return ""
    text = text.split()[0]

    # Already dd.mm.yyyy
    m = RE_DATE_FMT_DMY_DOT.fullmatch(text)
    if m:
        d, mo, y = m.groups()
        return f"{int(d):02d}.{int(mo):02d}.{y}"

    # yyyy-mm-dd or yyyy/mm/dd
    m = RE_DATE_FMT_YMD.fullmatch(text)
    if m:
        y, mo, d = m.groups()
        return f"{int(d):02d}.{int(mo):02d}.{y}"

    # dd/mm/yyyy - Indian/European interpretation by design.
    m = RE_DATE_FMT_DMY_SLASH.fullmatch(text)
    if m:
        d, mo, y = m.groups()
        return f"{int(d):02d}.{int(mo):02d}.{y}"

    return text


def output_paths(xlsx: str) -> dict:
    base = Path(xlsx)
    return {
        "xlsx": xlsx,
        "worklist": str(base.with_name(base.stem + "_WORKLIST.csv")),
        "status": str(base.with_name(base.stem + "_STATUS.csv")),
        "checkpoint": str(base.with_name(base.stem + "_CHECKPOINT.txt")),
        "sap_snapshot": str(base.with_name(base.stem + "_SAP_SNAPSHOT.csv")),
        "sap_reconcile": str(base.with_name(base.stem + "_SAP_RECONCILE_REPORT.xlsx")),
        "column_discovery": str(base.with_name(base.stem + "_SAP_COLUMN_DISCOVERY.csv")),
        "column_discovery_xlsx": str(base.with_name(base.stem + "_SAP_COLUMN_DISCOVERY.xlsx")),
        "validation": str(base.with_name(base.stem + "_VALIDATION_REPORT.csv")),
        "worklist_xlsx": str(base.with_name(base.stem + "_WORKLIST_VIEW.xlsx")),
        "validation_xlsx": str(base.with_name(base.stem + "_VALIDATION_REPORT.xlsx")),
        "status_xlsx": str(base.with_name(base.stem + "_STATUS_REPORT.xlsx")),
        "match": str(base.with_name(base.stem + "_MATCH_REPORT.csv")),
        "match_xlsx": str(base.with_name(base.stem + "_MATCH_REPORT.xlsx")),
        "vendor_dashboard": str(base.with_name(base.stem + "_VENDOR_DASHBOARD.xlsx")),
        "run_summary": str(base.with_name(base.stem + "_RUN_SUMMARY.xlsx")),
    }


def autosize_worksheet(ws, max_width=60):
    for col_cells in ws.columns:
        max_len = 0
        col_letter = openpyxl.utils.get_column_letter(col_cells[0].column)
        for cell in col_cells:
            val = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, min(len(val), max_width))
        ws.column_dimensions[col_letter].width = max(10, max_len + 2)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def write_sheets_to_run_report(xlsx_path: str, sheets_dict: dict):
    """Write or update multiple sheets in the consolidated RUN_REPORT workbook in a single open/save pass."""
    report_xlsx = str(Path(xlsx_path).with_name(Path(xlsx_path).stem + "_RUN_REPORT.xlsx"))
    
    if os.path.exists(report_xlsx):
        try:
            wb = openpyxl.load_workbook(report_xlsx)
        except Exception:
            wb = openpyxl.Workbook()
    else:
        wb = openpyxl.Workbook()
        
    for sheet_name, (fieldnames, records) in sheets_dict.items():
        title = sheet_name[:31]
        if title in wb.sheetnames:
            wb.remove(wb[title])
            
        ws = wb.create_sheet(title=title)
        ws.append(list(fieldnames))
        for rec in records:
            if isinstance(rec, dict):
                ws.append([rec.get(f, "") for f in fieldnames])
            elif isinstance(rec, (list, tuple)):
                ws.append(list(rec))
                
        autosize_worksheet(ws)
        
    if "Sheet" in wb.sheetnames and len(wb.sheetnames) > 1:
        wb.remove(wb["Sheet"])
        
    try:
        wb.save(report_xlsx)
        print(f"[RUN-REPORT] Updated sheets {list(sheets_dict.keys())} in {report_xlsx}")
    except PermissionError:
        print(f"[WARN] Could not save consolidated report because it is open: {report_xlsx}")
    finally:
        wb.close()


def write_to_run_report(xlsx_path: str, sheet_name: str, fieldnames, records):
    """Write or update a sheet in the consolidated RUN_REPORT workbook."""
    write_sheets_to_run_report(xlsx_path, {sheet_name: (fieldnames, records)})


def cleanup_old_reports(xlsx: str):
    paths = output_paths(xlsx)
    to_delete = [
        paths["worklist_xlsx"],
        paths["validation_xlsx"],
        paths["status_xlsx"],
        paths["match_xlsx"],
        paths["column_discovery_xlsx"],
        paths["vendor_dashboard"],
        paths["run_summary"],
        paths["sap_reconcile"]
    ]
    for p in to_delete:
        if os.path.exists(p):
            try:
                os.remove(p)
            except Exception:
                pass

    # If starting validate-export without a prior reconcile-sap step, clean up the consolidated report
    report_xlsx = str(Path(xlsx).with_name(Path(xlsx).stem + "_RUN_REPORT.xlsx"))
    if os.path.exists(report_xlsx):
        try:
            wb = openpyxl.load_workbook(report_xlsx)
            has_reconcile = "SAP Reconcile" in wb.sheetnames
            wb.close()
        except Exception:
            has_reconcile = False
        if not has_reconcile:
            try:
                os.remove(report_xlsx)
            except Exception:
                pass


def write_preflight_dashboards(paths, rows, validation_records, run_id, counts):
    vendor_summary = {}
    for item in rows:
        vendor = item.get("VendorName") or "UNKNOWN"
        rec = vendor_summary.setdefault(vendor, {"VendorName": vendor, "Rows": 0, "DuplicateKeyRows": 0, "ValidationErrors": 0, "PullInRows": 0, "DFT8Rows": 0, "DFT9Rows": 0})
        rec["Rows"] += 1
        if int(item.get("KeyTotal") or "1") > 1:
            rec["DuplicateKeyRows"] += 1
        code = dft_code(item.get("DFT", ""))
        if code == "8": rec["DFT8Rows"] += 1
        if code == "9": rec["DFT9Rows"] += 1
        if "pull in" in item.get("Remarks", "").lower():
            rec["PullInRows"] += 1
    errors = {r.get("ExcelRow") for r in validation_records if r.get("Severity") == "ERROR"}
    for item in rows:
        if item.get("ExcelRow") in errors:
            vendor_summary[item.get("VendorName") or "UNKNOWN"]["ValidationErrors"] += 1

    summary_rows = [
        {"Metric": "RunID", "Value": run_id},
        {"Metric": "Work Rows", "Value": len(rows)},
        {"Metric": "Validation Errors", "Value": sum(1 for r in validation_records if r.get("Severity") == "ERROR")},
        {"Metric": "Duplicate/Warn Rows", "Value": sum(1 for r in validation_records if r.get("Severity") == "WARN")},
        {"Metric": "Blank PO Skipped", "Value": counts.get("blank_po", 0)},
        {"Metric": "Hidden Excel Rows Skipped", "Value": counts.get("hidden_rows_skipped", 0)},
        {"Metric": "Master Rows Not In SAP Snapshot Skipped", "Value": counts.get("skipped_not_in_snapshot", 0)},
        {"Metric": "Already Done Skipped", "Value": counts.get("skipped_done", 0)},
        {"Metric": "Non-Failed Skipped", "Value": counts.get("skipped_not_failed", 0)},
    ]

    write_sheets_to_run_report(paths["xlsx"], {
        "Vendor Dashboard": (["VendorName", "Rows", "DuplicateKeyRows", "ValidationErrors", "PullInRows", "DFT8Rows", "DFT9Rows"], list(vendor_summary.values())),
        "Run Summary": (["Metric", "Value"], summary_rows)
    })


def load_workbook_rw(path: str):
    if not os.path.exists(path):
        sys.exit(f"[FATAL] File not found: {path}")
    try:
        return openpyxl.load_workbook(path)
    except PermissionError:
        sys.exit("[FATAL] XLSX is locked. Close it in LibreOffice Calc/Excel and retry.")
    except Exception as exc:
        sys.exit(f"[FATAL] Could not open XLSX: {exc}")


def normalize_digits(value):
    text = RE_NON_DIGITS.sub("", sanitize(value))
    return text.lstrip("0") or ("0" if text else "")


def normalize_material(value):
    text = sanitize(value).upper().replace(" ", "")
    if text.isdigit():
        return text.lstrip("0") or "0"
    return text


def normalize_qty(value):
    text = sanitize(value).upper().strip()
    if not text:
        return ""
    # Remove units/text while keeping digits, comma, dot, minus.
    text = RE_QTY_CLEAN.sub("", text)
    if not text:
        return ""
    # Handle decimal comma if no dot exists: 6,00 -> 6.00
    if "," in text and "." not in text:
        parts = text.split(",")
        if len(parts[-1]) in (1, 2, 3):
            text = "".join(parts[:-1]) + "." + parts[-1]
        else:
            text = text.replace(",", "")
    else:
        text = text.replace(",", "")
    try:
        num = float(text)
        return str(int(num)) if num.is_integer() else ("%.6f" % num).rstrip("0").rstrip(".")
    except ValueError:
        return normalize_digits(text)


def normalize_date_key(value):
    d = parse_date(value)
    return d.strftime("%Y%m%d") if d else ""


def make_match_key(item):
    return "|".join([
        normalize_digits(item.get("PO", "")),
        normalize_material(item.get("MaterialCode", "")),
        normalize_digits(item.get("Line", "")),
        normalize_date_key(item.get("SchedDate", "")),
        normalize_qty(item.get("Qty", "")),
    ])


def parse_int(value):
    text = sanitize(value)
    if not text:
        return None
    try:
        return int(float(text))
    except ValueError:
        return None


def dft_code(value):
    text = sanitize(value).strip()
    if not text:
        return ""
    m = RE_DFT_CODE.match(text)
    return m.group(1) if m else text


def is_manual_status_row(remarks, dft):
    code = dft_code(dft)
    if code in MANUAL_DFT_CODES:
        return True
    r = sanitize(remarks).lower()
    return any(k in r for k in MANUAL_REMARK_KEYWORDS)


def is_excel_row_hidden(ws, row_num):
    return bool(ws.row_dimensions[row_num].hidden)


def apply_excel_automation(ws):
    """Apply business automation to the workbook before worklist export.

    User input model:
      - Remarks column may contain only supplier committed dispatch date.
      - Supplier Lead Time column contains days from committed dispatch to expected receipt.
      - Expected Receive Date is stored in existing K / expected-date column.
      - DFT is auto 8/9 unless row is a manual-status case.
    """
    changes = []

    lead_col, created_lead_col = ensure_header_col(ws, SUPPLIER_LEAD_TIME_HEADER)
    if created_lead_col:
        changes.append(f"Created helper column '{SUPPLIER_LEAD_TIME_HEADER}' at {openpyxl.utils.get_column_letter(lead_col)}")

    committed_col, created_committed_col = ensure_header_col(ws, COMMITTED_DATE_HEADER)
    if created_committed_col:
        changes.append(f"Created helper column '{COMMITTED_DATE_HEADER}' at {openpyxl.utils.get_column_letter(committed_col)}")

    mode_col, created_mode_col = ensure_header_col(ws, AUTOMATION_MODE_HEADER)
    if created_mode_col:
        changes.append(f"Created helper column '{AUTOMATION_MODE_HEADER}' at {openpyxl.utils.get_column_letter(mode_col)}")

    for row_num in range(DATA_START_ROW, ws.max_row + 1):
        if is_excel_row_hidden(ws, row_num):
            continue
        po = sanitize(ws.cell(row_num, COLS["po"]).value)
        if not po or po == "0":
            continue

        remarks_cell = ws.cell(row_num, COLS["remarks"])
        exp_cell = ws.cell(row_num, COLS["exp_date"])
        dft_cell = ws.cell(row_num, COLS["dft"])
        sched_date = parse_date(ws.cell(row_num, COLS["sched_date"]).value)
        remarks_before = sanitize(remarks_cell.value)
        dft_before = sanitize(dft_cell.value)
        mode = sanitize(ws.cell(row_num, mode_col).value).upper()
        if mode == "":
            mode = "AUTO"
            ws.cell(row_num, mode_col).value = mode
            changes.append(f"Row {row_num}: Automation Mode defaulted to AUTO")

        if mode == "SKIP":
            continue
        if mode == "MANUAL" or is_manual_status_row(remarks_before, dft_before):
            if sanitize(ws.cell(row_num, mode_col).value).upper() != "MANUAL":
                ws.cell(row_num, mode_col).value = "MANUAL"
                changes.append(f"Row {row_num}: Automation Mode set to MANUAL")
            continue

        committed_date = parse_date(ws.cell(row_num, committed_col).value)
        remarks_has_date_only = parse_date_only(remarks_before)
        if committed_date is None:
            # Backward compatible and safe: only if Remarks contains a date alone, no other text/content.
            committed_date = remarks_has_date_only
            if committed_date is not None:
                ws.cell(row_num, committed_col).value = committed_date.strftime("%d.%m.%Y")
                changes.append(f"Row {row_num}: Supplier Committed Date captured from date-only Remarks")

        if committed_date is None:
            # No committed date available remains a manual case.
            continue

        lead_days = parse_int(ws.cell(row_num, lead_col).value)

        dispatch_remark = AUTO_REMARK_TEMPLATE.format(date=committed_date.strftime("%d.%m.%Y"))
        # Only auto-change Remarks if it is blank, date-only, or already system-generated.
        # Do not overwrite manual remark content.
        can_auto_update_remark = (remarks_before == "" or remarks_has_date_only is not None or is_auto_generated_remark(remarks_before))
        if can_auto_update_remark and remarks_before != dispatch_remark:
            remarks_cell.value = dispatch_remark
            changes.append(f"Row {row_num}: Remarks auto-formatted")

        if lead_days is None:
            # Remark can still be automated, but receive date/DFT need supplier lead time.
            continue

        expected_receive = committed_date + dt.timedelta(days=lead_days)
        if parse_date(exp_cell.value) != expected_receive:
            exp_cell.value = expected_receive.strftime("%d.%m.%Y")
            changes.append(f"Row {row_num}: Expected Receive Date calculated")

        if sched_date is None:
            continue

        if expected_receive <= sched_date:
            if dft_code(dft_cell.value) != "8":
                dft_cell.value = "8"
                changes.append(f"Row {row_num}: DFT set to 8")
        else:
            if dft_code(dft_cell.value) != "9":
                dft_cell.value = "9"
                changes.append(f"Row {row_num}: DFT set to 9")
            pull_in_remark = PULL_IN_REMARK_TEMPLATE.format(date=committed_date.strftime("%d/%m/%Y"))
            if can_auto_update_remark and sanitize(remarks_cell.value) != pull_in_remark:
                remarks_cell.value = pull_in_remark
                changes.append(f"Row {row_num}: Pull-in remark generated with committed date")

    return changes


def validate_headers(ws):
    errors = []
    for col, expected in HEADER_CHECKS:
        actual = sanitize(ws.cell(HEADER_ROW, col).value)
        if expected.lower() not in actual.lower():
            letter = openpyxl.utils.get_column_letter(col)
            errors.append(f"Col {letter}: expected header containing '{expected}', found '{actual}'")
    return errors


def validate_row(row_data):
    errors = []
    po = row_data["PO"]
    if not po:
        return errors

    status_in = row_data["StatusIn"].strip().lower()
    if status_in == DONE_TEXT.lower():
        return errors

    line = row_data.get("Line", "")
    material = row_data.get("MaterialCode", "")
    sched = row_data.get("SchedDate", "")
    qty = row_data.get("Qty", "")
    pernr = row_data["PERNR"]
    cummode = row_data["CUMMODE"].upper()
    exp = row_data["ExpDate"]
    remarks = row_data["Remarks"]
    dft = row_data["DFT"]

    if not normalize_digits(po):
        errors.append("PO Number missing/invalid")
    if not normalize_digits(line):
        errors.append("PO Line Item missing/invalid")
    if not material:
        errors.append("Material Code missing")

    if not sched or not parse_date(sched):
        errors.append(f"Schedule Date missing/invalid [{sched}]")
    if not qty or normalize_qty(qty) == "":
        errors.append(f"Pending Qty missing/invalid [{qty}]")
    else:
        try:
            if float(sanitize(qty).replace(",", "")) <= 0:
                errors.append(f"Pending Qty must be > 0 [{qty}]")
        except ValueError:
            errors.append(f"Pending Qty invalid [{qty}]")

    if exp:
        parsed_exp = parse_date(exp)
        if not parsed_exp:
            errors.append(f"Expected Receive Date invalid [{exp}]")
        elif parsed_exp < dt.date.today():
            errors.append(f"Expected Delivery Date (EXPDATE) [{exp}] is before today's date [{dt.date.today().strftime('%d.%m.%Y')}]")

    if not pernr:
        errors.append("PERNR missing")
    elif not pernr.isdigit() or len(pernr) != 6:
        errors.append(f"invalid PERNR '{pernr}' - must be exactly 6 digits")

    if not cummode:
        errors.append("CUMMODE missing")
    elif cummode not in VALID_CUMMODES:
        errors.append(f"invalid CUMMODE '{cummode}' - allowed: EMAIL, PHONE, VISIT, SANYOG")

    if not exp and not remarks and not dft:
        errors.append("Expected Date / Remarks / DFT are all blank")

    if remarks.startswith("Material is expected to dispatch by") and not exp:
        errors.append("Expected Receive Date missing - enter Supplier Lead Time for committed-date rows")

    code = dft_code(dft)
    if dft and (not code.isdigit() or int(code) < 1 or int(code) > 31):
        errors.append(f"DFT invalid [{dft}]")

    if exp and (remarks.startswith("Material is expected to dispatch by") or remarks.startswith("Request supplier to pull in")) and code not in {"8", "9"}:
        errors.append("Auto committed-date rows must have DFT 8 or 9")

    return errors


def row_to_workitem(ws, row_num, run_id="", pernr_override="", cummode_override=""):
    def c(key):
        return ws.cell(row_num, COLS[key]).value

    item = {
        "RunID": run_id,
        "ExcelRow": str(row_num),
        "VendorName": sanitize(c("vendor_name")),
        "PO": sanitize(c("po")),
        "Line": sanitize(c("line")),
        "MaterialCode": sanitize(c("material_code")),
        "SchedDate": format_date(c("sched_date")),
        "Qty": sanitize(c("qty")),
        "MatchKey": "",
        "KeyOccurrence": "",
        "KeyTotal": "",
        "SAPRow": "",
        "ExpDate": format_date(c("exp_date")),
        "Remarks": sanitize(c("remarks")),
        "DFT": sanitize(c("dft")),
        "PERNR": sanitize(pernr_override) if pernr_override else sanitize(c("pernr")),
        "CUMMODE": (sanitize(cummode_override) if cummode_override else sanitize(c("cummode"))).upper(),
        "StatusIn": sanitize(c("status_in")),
        "StatusOut": sanitize(c("status_out")),
    }
    item["MatchKey"] = make_match_key(item)
    return item


def backup_workbook(xlsx: str, run_id: str):
    src = Path(xlsx)
    backup_dir = src.with_name("Backups")
    backup_dir.mkdir(exist_ok=True)
    dst = backup_dir / f"{src.stem}_{run_id}.xlsx"
    try:
        shutil.copy2(src, dst)
        print(f"[BACKUP] {dst}")
    except Exception as exc:
        print(f"[WARN] Could not create backup: {exc}")


def validate_runtime_inputs(pernr: str, cummode: str):
    errs = []
    if not pernr or not pernr.isdigit() or len(pernr) != 6:
        errs.append(f"Runtime PERNR must be exactly 6 digits (got: {pernr})")
    if not cummode or cummode.upper() not in VALID_CUMMODES:
        errs.append(f"Runtime Communication Mode invalid [{cummode}] - allowed: EMAIL, PHONE, VISIT, SANYOG")
    return errs


def load_sap_snapshot(snapshot_csv: str):
    if not os.path.exists(snapshot_csv):
        return []
    with open(snapshot_csv, "r", newline="", encoding="utf-8-sig", errors="replace") as f:
        return list(csv.DictReader(f))


def snapshot_item_to_key(item):
    return "|".join([
        normalize_digits(item.get("PO", "")),
        normalize_material(item.get("MaterialCode", "")),
        normalize_digits(item.get("Line", "")),
        normalize_date_key(item.get("SchedDate", "")),
        normalize_qty(item.get("Qty", "")),
    ])


def is_relevant_sap_snapshot_row(item):
    return all([
        normalize_digits(item.get("PO", "")),
        normalize_material(item.get("MaterialCode", "")),
        normalize_digits(item.get("Line", "")),
        normalize_date_key(item.get("SchedDate", "")),
        normalize_qty(item.get("Qty", "")),
    ])


def reconcile_sap_snapshot(xlsx: str, snapshot_csv: str = "", run_id: str = "") -> int:
    run_id = run_id or dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    paths = output_paths(xlsx)
    snapshot_csv = snapshot_csv or paths["sap_snapshot"]

    # Delete old consolidated report at start of reconciliation run
    report_xlsx = str(Path(xlsx).with_name(Path(xlsx).stem + "_RUN_REPORT.xlsx"))
    if os.path.exists(report_xlsx):
        try:
            os.remove(report_xlsx)
        except Exception:
            pass

    rows = [r for r in load_sap_snapshot(snapshot_csv) if is_relevant_sap_snapshot_row(r)]
    if not rows:
        print(f"[WARN] SAP snapshot empty or not found: {snapshot_csv}")
        return 0

    backup_workbook(xlsx, run_id)
    wb = load_workbook_rw(xlsx)
    if SHEET_NAME not in wb.sheetnames:
        sys.exit(f"[FATAL] Sheet not found: {SHEET_NAME}")
    ws = wb[SHEET_NAME]

    # Existing master rows by duplicate-aware key occurrence.
    # IMPORTANT: user-maintained fields are preserved for existing rows:
    # K Expected Date, L Remarks, M DFT, O/P runtime fields, Q/R/S status.
    existing_by_key = {}
    existing_po = set()
    existing_po_mat = set()
    existing_po_mat_line = set()
    existing_po_mat_line_date = set()
    for row_num in range(DATA_START_ROW, ws.max_row + 1):
        item = row_to_workitem(ws, row_num, run_id=run_id)
        if not item["PO"] or item["PO"] == "0":
            continue
        po_n = normalize_digits(item.get("PO", ""))
        mat_n = normalize_material(item.get("MaterialCode", ""))
        line_n = normalize_digits(item.get("Line", ""))
        date_n = normalize_date_key(item.get("SchedDate", ""))
        existing_po.add(po_n)
        existing_po_mat.add((po_n, mat_n))
        existing_po_mat_line.add((po_n, mat_n, line_n))
        existing_po_mat_line_date.add((po_n, mat_n, line_n, date_n))
        existing_by_key.setdefault(item["MatchKey"], []).append(row_num)

    snap_seen = Counter()
    new_info = 0
    matched_existing = 0
    unchanged = 0
    report = []

    for snap in rows:
        key = snapshot_item_to_key(snap)
        snap_seen[key] += 1
        occ = snap_seen[key]
        existing_rows = existing_by_key.get(key, [])
        action = "UNCHANGED"
        target_row = None

        if occ <= len(existing_rows):
            target_row = existing_rows[occ - 1]
            # Existing master row found. Do NOT modify master data here.
            # User remarks/updates and even reference fields are preserved.
            action = "FOUND_IN_MASTER_NO_CHANGE"
            matched_existing += 1
            unchanged += 1
        else:
            po_n = normalize_digits(snap.get("PO", ""))
            mat_n = normalize_material(snap.get("MaterialCode", ""))
            line_n = normalize_digits(snap.get("Line", ""))
            date_n = normalize_date_key(snap.get("SchedDate", ""))
            if po_n not in existing_po:
                action = "NEW_IN_SAP_PO_REVIEW_ONLY"
            elif (po_n, mat_n) not in existing_po_mat:
                action = "NEW_IN_SAP_MATERIAL_REVIEW_ONLY"
            elif (po_n, mat_n, line_n) not in existing_po_mat_line:
                action = "NEW_IN_SAP_PO_LINE_REVIEW_ONLY"
            elif (po_n, mat_n, line_n, date_n) not in existing_po_mat_line_date:
                action = "NEW_IN_SAP_SCHEDULE_LINE_REVIEW_ONLY"
            else:
                action = "NEW_IN_SAP_QTY_OR_DUPLICATE_REVIEW_ONLY"

            # IMPORTANT: Do not append or modify master Excel.
            # New SAP rows are only reported to the user for manual review/import decision.
            target_row = ""
            new_info += 1

        report.append({
            "RunID": run_id,
            "Action": action,
            "ExcelRow": target_row,
            "SAPRow": snap.get("SAPRow", ""),
            "VendorName": snap.get("VendorName", ""),
            "PO": snap.get("PO", ""),
            "Line": snap.get("Line", ""),
            "MaterialCode": snap.get("MaterialCode", ""),
            "SchedDate": format_date(snap.get("SchedDate", "")),
            "Qty": snap.get("Qty", ""),
            "MatchKey": key,
            "KeyOccurrence": occ,
            "PreservedUserFields": "YES" if action == "FOUND_IN_MASTER_NO_CHANGE" else "NOT_IMPORTED_REVIEW_ONLY",
        })

    # Reconcile is now report-only: no SAP snapshot rows are appended and no existing master data is changed.
    wb.close()

    write_to_run_report(xlsx, "SAP Reconcile", ["RunID", "Action", "ExcelRow", "SAPRow", "VendorName", "PO", "Line", "MaterialCode", "SchedDate", "Qty", "MatchKey", "KeyOccurrence", "PreservedUserFields"], report)
    print(f"[RECONCILE] Snapshot rows={len(rows)} found_in_master={matched_existing} new_in_sap_review_only={new_info} master_changed=NO")
    return 0


def saprow_by_key_occurrence(snapshot_csv: str):
    mapping = {}
    rows = [r for r in load_sap_snapshot(snapshot_csv) if is_relevant_sap_snapshot_row(r)]
    seen = Counter()
    for r in rows:
        key = snapshot_item_to_key(r)
        seen[key] += 1
        mapping[(key, seen[key])] = sanitize(r.get("SAPRow"))
    return mapping


def validate_export(xlsx: str, worklist, write_validation_errors: bool, pernr: str = "", cummode: str = "", run_id: str = "", rerun_mode: str = "ALL") -> int:
    run_id = run_id or dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    pernr = sanitize(pernr)
    cummode = sanitize(cummode).upper()
    runtime_errors = validate_runtime_inputs(pernr, cummode)
    if runtime_errors:
        print("\n[RUNTIME INPUT VALIDATION FAILED]")
        for err in runtime_errors:
            print(" ", err)
        return 1

    backup_workbook(xlsx, run_id)
    wb = load_workbook_rw(xlsx)
    if SHEET_NAME not in wb.sheetnames:
        sys.exit(f"[FATAL] Sheet not found: {SHEET_NAME}")
    ws = wb[SHEET_NAME]

    print("\nSAP Tracker Validation/Export v5 LibreOffice-safe")
    print("-" * 60)
    print(f"[RUN]  {run_id}")
    print(f"[XLSX] {xlsx}")
    print(f"[SHEET] {SHEET_NAME}")
    print(f"[READ] max_row={ws.max_row}, max_column={ws.max_column}")

    header_errors = validate_headers(ws)
    if header_errors:
        print("\n[HEADER VALIDATION FAILED]")
        for err in header_errors:
            print(" ", err)
        wb.close()
        return 1
    print("[HDR] Headers OK")

    automation_changes = apply_excel_automation(ws)
    if automation_changes:
        print(f"[AUTO] Applied Excel automation changes: {len(automation_changes)}")
        for msg in automation_changes[:50]:
            print(" ", msg)
        if len(automation_changes) > 50:
            print(f"  ... plus {len(automation_changes) - 50} more")
        try:
            wb.save(xlsx)
        except PermissionError:
            wb.close()
            sys.exit("[FATAL] Could not save automation changes. Close XLSX in LibreOffice and retry.")
    else:
        print("[AUTO] No Excel automation changes needed")

    paths = output_paths(xlsx)
    worklist = worklist or paths["worklist"]

    rows = []
    validation_errors = []
    skipped_done = 0
    skipped_not_failed = 0
    blank_po = 0
    hidden_rows_skipped = 0
    skipped_not_in_snapshot = 0
    snapshot_warning_records = []
    runtime_populated = 0
    rerun_mode = sanitize(rerun_mode).upper() or "ALL"

    for row_num in range(DATA_START_ROW, ws.max_row + 1):
        if is_excel_row_hidden(ws, row_num):
            hidden_rows_skipped += 1
            continue
        item = row_to_workitem(ws, row_num, run_id=run_id, pernr_override=pernr, cummode_override=cummode)
        if not item["PO"] or item["PO"] == "0":
            blank_po += 1
            continue
        status_out_norm = item.get("StatusOut", "").strip().lower()
        if item["StatusIn"].strip().lower() == DONE_TEXT.lower() or status_out_norm == DONE_TEXT.lower():
            skipped_done += 1
            continue
        if rerun_mode == "FAILED" and status_out_norm not in {"sap error", "validation error", "missing data"}:
            skipped_not_failed += 1
            continue
        # Populate runtime PERNR/CUMMODE into workbook columns O/P for auditability.
        if sanitize(ws.cell(row_num, COLS["pernr"]).value) != pernr:
            ws.cell(row_num, COLS["pernr"]).value = pernr
            runtime_populated += 1
        if sanitize(ws.cell(row_num, COLS["cummode"]).value).upper() != cummode:
            ws.cell(row_num, COLS["cummode"]).value = cummode
            runtime_populated += 1
        rows.append(item)

    # Duplicate audit: duplicates are allowed, but counted and sequenced.
    saprow_map = saprow_by_key_occurrence(paths["sap_snapshot"])
    snapshot_exists = os.path.exists(paths["sap_snapshot"])
    key_counts = Counter(item["MatchKey"] for item in rows)
    key_seen = Counter()
    active_rows = []
    for item in rows:
        key = item["MatchKey"]
        key_seen[key] += 1
        item["KeyOccurrence"] = str(key_seen[key])
        item["KeyTotal"] = str(key_counts[key])
        item["SAPRow"] = saprow_map.get((key, key_seen[key]), "")

        # If master row is not present in the current SAP snapshot, do NOT block the run.
        # Report it as WARN and skip it from the SAP worklist so valid matched master rows can update.
        if snapshot_exists and not item["SAPRow"]:
            skipped_not_in_snapshot += 1
            snapshot_warning_records.append({
                "RunID": item.get("RunID", ""),
                "ExcelRow": item.get("ExcelRow", ""),
                "VendorName": item.get("VendorName", ""),
                "PO": item.get("PO", ""),
                "Line": item.get("Line", ""),
                "MaterialCode": item.get("MaterialCode", ""),
                "SchedDate": item.get("SchedDate", ""),
                "Qty": item.get("Qty", ""),
                "MatchKey": item.get("MatchKey", ""),
                "KeyOccurrence": item.get("KeyOccurrence", ""),
                "KeyTotal": item.get("KeyTotal", ""),
                "Severity": "WARN",
                "Message": "Master row not found in current SAP snapshot - skipped from this run; master data not changed",
            })
            continue

        active_rows.append(item)
        errs = validate_row(item)
        if key_counts[key] > 1:
            # Audit warning only. VBS uses next unused SAP match for duplicate keys.
            pass
        if errs:
            row_num = int(item["ExcelRow"])
            validation_errors.append((row_num, item["PO"], " | ".join(errs)))
            if write_validation_errors:
                ws.cell(row_num, COLS["status_out"]).value = "Validation Error"
                ws.cell(row_num, COLS["err_msg"]).value = " | ".join(errs)
                ws.cell(row_num, COLS["upd_time"]).value = dt.datetime.now().strftime("%d.%m.%Y %H:%M:%S")

    rows = active_rows

    if runtime_populated:
        print(f"[AUTO] Runtime PERNR/CUMMODE populated in workbook cells: {runtime_populated}")

    if runtime_populated or (write_validation_errors and validation_errors):
        try:
            wb.save(xlsx)
        except PermissionError:
            wb.close()
            sys.exit("[FATAL] Could not save workbook updates. Close XLSX in LibreOffice and retry.")
    wb.close()

    with open(worklist, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=WORKLIST_FIELDS, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)

    validation_path = paths["validation"]
    validation_fieldnames = ["RunID", "ExcelRow", "VendorName", "PO", "Line", "MaterialCode", "SchedDate", "Qty", "MatchKey", "KeyOccurrence", "KeyTotal", "Severity", "Message"]
    validation_records = []
    error_by_row = {str(r): msg for r, _po, msg in validation_errors}
    for item in rows:
        severity = "ERROR" if item["ExcelRow"] in error_by_row else ("WARN" if int(item["KeyTotal"] or "1") > 1 else "OK")
        message = error_by_row.get(item["ExcelRow"], "Duplicate key - sequenced" if severity == "WARN" else "")
        rec = {k: item.get(k, "") for k in validation_fieldnames if k not in {"Severity", "Message"}}
        rec["Severity"] = severity
        rec["Message"] = message
        validation_records.append(rec)

    validation_records.extend(snapshot_warning_records)

    with open(validation_path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=validation_fieldnames)
        writer.writeheader()
        writer.writerows(validation_records)

    # Clean up any old individual reports and run report
    cleanup_old_reports(xlsx)

    # Calculate preflight dashboards
    vendor_summary = {}
    for item in rows:
        vendor = item.get("VendorName") or "UNKNOWN"
        rec = vendor_summary.setdefault(vendor, {"VendorName": vendor, "Rows": 0, "DuplicateKeyRows": 0, "ValidationErrors": 0, "PullInRows": 0, "DFT8Rows": 0, "DFT9Rows": 0})
        rec["Rows"] += 1
        if int(item.get("KeyTotal") or "1") > 1:
            rec["DuplicateKeyRows"] += 1
        code = dft_code(item.get("DFT", ""))
        if code == "8": rec["DFT8Rows"] += 1
        if code == "9": rec["DFT9Rows"] += 1
        if "pull in" in item.get("Remarks", "").lower():
            rec["PullInRows"] += 1
    errors = {r.get("ExcelRow") for r in validation_records if r.get("Severity") == "ERROR"}
    for item in rows:
        if item.get("ExcelRow") in errors:
            vendor_summary[item.get("VendorName") or "UNKNOWN"]["ValidationErrors"] += 1

    counts = {"blank_po": blank_po, "hidden_rows_skipped": hidden_rows_skipped, "skipped_not_in_snapshot": skipped_not_in_snapshot, "skipped_done": skipped_done, "skipped_not_failed": skipped_not_failed}
    summary_rows = [
        {"Metric": "RunID", "Value": run_id},
        {"Metric": "Work Rows", "Value": len(rows)},
        {"Metric": "Validation Errors", "Value": sum(1 for r in validation_records if r.get("Severity") == "ERROR")},
        {"Metric": "Duplicate/Warn Rows", "Value": sum(1 for r in validation_records if r.get("Severity") == "WARN")},
        {"Metric": "Blank PO Skipped", "Value": counts.get("blank_po", 0)},
        {"Metric": "Hidden Excel Rows Skipped", "Value": counts.get("hidden_rows_skipped", 0)},
        {"Metric": "Master Rows Not In SAP Snapshot Skipped", "Value": counts.get("skipped_not_in_snapshot", 0)},
        {"Metric": "Already Done Skipped", "Value": counts.get("skipped_done", 0)},
        {"Metric": "Non-Failed Skipped", "Value": counts.get("skipped_not_failed", 0)},
    ]

    # Batch write all 4 pre-flight sheets directly to RUN_REPORT.xlsx in a single open/save pass
    write_sheets_to_run_report(xlsx, {
        "Worklist": (WORKLIST_FIELDS, rows),
        "Validation": (validation_fieldnames, validation_records),
        "Vendor Dashboard": (["VendorName", "Rows", "DuplicateKeyRows", "ValidationErrors", "PullInRows", "DFT8Rows", "DFT9Rows"], list(vendor_summary.values())),
        "Run Summary": (["Metric", "Value"], summary_rows)
    })

    print(f"[EXPORT] Worklist written: {worklist}")
    print(f"[REPORT] Validation report: {validation_path}")
    print(f"[INFO] Work rows: {len(rows)} | hidden rows skipped: {hidden_rows_skipped} | not in SAP snapshot skipped: {skipped_not_in_snapshot} | blank PO skipped: {blank_po} | already-done skipped: {skipped_done} | non-failed skipped: {skipped_not_failed}")

    if validation_errors:
        print("\n[VALIDATION FAILED]")
        for row_num, po, msg in validation_errors[:200]:
            print(f"  Row {row_num} (PO {po}): {msg}")
        if len(validation_errors) > 200:
            print(f"  ... plus {len(validation_errors) - 200} more")
        return 1

    print("\n[OK] All validations passed")
    return 0


def convert_reports(xlsx: str) -> int:
    paths = output_paths(xlsx)
    sheets_to_write = {}
    
    # Collect 'Status' sheet
    if os.path.exists(paths["status"]):
        with open(paths["status"], "r", newline="", encoding="utf-8-sig", errors="replace") as f:
            reader = csv.DictReader(f)
            status_fieldnames = reader.fieldnames or STATUS_FIELDS
            sheets_to_write["Status"] = (status_fieldnames, list(reader))
            
    # Collect 'SAP Match' sheet
    if os.path.exists(paths["match"]):
        with open(paths["match"], "r", newline="", encoding="utf-8-sig", errors="replace") as f:
            reader = csv.DictReader(f)
            match_fieldnames = reader.fieldnames or ["RunID", "ExcelRow", "VendorName", "PO", "Line", "MaterialCode", "SchedDate", "Qty", "MatchKey", "KeyOccurrence", "KeyTotal", "SAPRow", "Status", "Message"]
            sheets_to_write["SAP Match"] = (match_fieldnames, list(reader))
            
    # Collect 'SAP Columns' sheet
    if os.path.exists(paths["column_discovery"]):
        with open(paths["column_discovery"], "r", newline="", encoding="utf-8-sig", errors="replace") as f:
            reader = csv.DictReader(f)
            col_fieldnames = reader.fieldnames or ["Index", "Name", "TechName", "Tooltip"]
            sheets_to_write["SAP Columns"] = (col_fieldnames, list(reader))
            
    if sheets_to_write:
        write_sheets_to_run_report(xlsx, sheets_to_write)

    print("[OK] Report conversion complete")
    return 0


def apply_status(xlsx: str, status_csv) -> int:
    paths = output_paths(xlsx)
    status_csv = status_csv or paths["status"]

    if not os.path.exists(status_csv):
        sys.exit(f"[FATAL] Status CSV not found: {status_csv}")

    wb = load_workbook_rw(xlsx)
    if SHEET_NAME not in wb.sheetnames:
        sys.exit(f"[FATAL] Sheet not found: {SHEET_NAME}")
    ws = wb[SHEET_NAME]

    updates = 0
    skipped = 0
    with open(status_csv, "r", newline="", encoding="utf-8-sig", errors="replace") as f:
        reader = csv.DictReader(f)
        for rec in reader:
            excel_row = sanitize(rec.get("ExcelRow"))
            if not excel_row.isdigit():
                skipped += 1
                continue
            row_num = int(excel_row)
            if row_num < DATA_START_ROW or row_num > ws.max_row:
                skipped += 1
                continue
            ws.cell(row_num, COLS["status_out"]).value = sanitize(rec.get("Status"))
            ws.cell(row_num, COLS["err_msg"]).value = sanitize(rec.get("ErrorRemarks"))
            ws.cell(row_num, COLS["upd_time"]).value = sanitize(rec.get("Timestamp"))
            updates += 1

    try:
        wb.save(xlsx)
    except PermissionError:
        wb.close()
        sys.exit("[FATAL] Could not save XLSX. Close it in LibreOffice Calc and retry.")
    wb.close()

    # Update 'Status' and 'SAP Match' sheets in consolidated RUN_REPORT.xlsx
    if os.path.exists(status_csv):
        with open(status_csv, "r", newline="", encoding="utf-8-sig", errors="replace") as f:
            reader = csv.DictReader(f)
            status_fieldnames = reader.fieldnames or STATUS_FIELDS
            write_to_run_report(xlsx, "Status", status_fieldnames, list(reader))
            
    if os.path.exists(paths["match"]):
        with open(paths["match"], "r", newline="", encoding="utf-8-sig", errors="replace") as f:
            reader = csv.DictReader(f)
            match_fieldnames = reader.fieldnames or ["RunID", "ExcelRow", "VendorName", "PO", "Line", "MaterialCode", "SchedDate", "Qty", "MatchKey", "KeyOccurrence", "KeyTotal", "SAPRow", "Status", "Message"]
            write_to_run_report(xlsx, "SAP Match", match_fieldnames, list(reader))

    print(f"[APPLY] Updated workbook rows: {updates}")
    if skipped:
        print(f"[APPLY] Skipped non-data/status rows: {skipped}")
    print(f"[OK] XLSX tracker updated: {xlsx}")
    return 0


def apply_single_status(xlsx: str, row_num: int, status: str, err_msg: str, upd_time: str) -> int:
    wb = load_workbook_rw(xlsx)
    if SHEET_NAME not in wb.sheetnames:
        sys.exit(f"[FATAL] Sheet not found: {SHEET_NAME}")
    ws = wb[SHEET_NAME]

    if row_num < DATA_START_ROW or row_num > ws.max_row:
        wb.close()
        sys.exit(f"[FATAL] Row number {row_num} out of bounds")

    ws.cell(row_num, COLS["status_out"]).value = sanitize(status)
    ws.cell(row_num, COLS["err_msg"]).value = sanitize(err_msg)
    ws.cell(row_num, COLS["upd_time"]).value = sanitize(upd_time)

    try:
        wb.save(xlsx)
    except PermissionError:
        wb.close()
        sys.exit(f"[FATAL] Could not save XLSX. Close it in LibreOffice Calc/Excel and retry.")
    wb.close()
    return 0


def main():
    parser = argparse.ArgumentParser(description="LibreOffice-safe SAP Daily Input validator/exporter/status applier")
    parser.add_argument("--xlsx", default=DEFAULT_XLSX)
    parser.add_argument("--mode", choices=["reconcile-sap", "validate-export", "apply-status", "convert-reports", "apply-single-status"], default="validate-export")
    parser.add_argument("--worklist", default=None)
    parser.add_argument("--status-csv", default=None)
    parser.add_argument("--sap-snapshot", default=None)
    parser.add_argument("--pernr", default="", help="Runtime Employee ID to apply to all processed rows")
    parser.add_argument("--cummode", default="", help="Runtime communication mode: EMAIL/PHONE/VISIT/SANYOG")
    parser.add_argument("--run-id", default="", help="Unique run id passed from BAT")
    parser.add_argument("--rerun", choices=["ALL", "FAILED"], default="ALL", help="ALL rows or only FAILED rows")
    parser.add_argument("--write-validation-errors", action="store_true")
    parser.add_argument("--row", type=int, help="Excel row number to update")
    parser.add_argument("--status", default="", help="Status to write")
    parser.add_argument("--err-msg", default="", help="Error message to write")
    parser.add_argument("--upd-time", default="", help="Update timestamp to write")
    args = parser.parse_args()

    if args.mode == "reconcile-sap":
        raise SystemExit(reconcile_sap_snapshot(args.xlsx, args.sap_snapshot, args.run_id))
    if args.mode == "validate-export":
        raise SystemExit(validate_export(args.xlsx, args.worklist, args.write_validation_errors, args.pernr, args.cummode, args.run_id, args.rerun))
    if args.mode == "convert-reports":
        raise SystemExit(convert_reports(args.xlsx))
    if args.mode == "apply-single-status":
        raise SystemExit(apply_single_status(args.xlsx, args.row, args.status, args.err_msg, args.upd_time))
    raise SystemExit(apply_status(args.xlsx, args.status_csv))


if __name__ == "__main__":
    main()
